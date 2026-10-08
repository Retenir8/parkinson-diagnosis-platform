#!/usr/bin/env python3
"""Segment a RealSense bag file and export pose keypoints, gait metrics, and videos.

This script is designed for Windows machines with Intel RealSense Python
bindings installed. It keeps the pyrealsense2 import inside the bag-processing
path so that --help and --self-test can run on platforms where pyrealsense2 is
not available.
"""

from __future__ import annotations

import argparse
import csv
import math
import os
import re
import statistics
import sys
import tempfile
import urllib.request
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

import numpy as np


LANDMARK_NAMES = [
    "nose",
    "left_eye_inner",
    "left_eye",
    "left_eye_outer",
    "right_eye_inner",
    "right_eye",
    "right_eye_outer",
    "left_ear",
    "right_ear",
    "mouth_left",
    "mouth_right",
    "left_shoulder",
    "right_shoulder",
    "left_elbow",
    "right_elbow",
    "left_wrist",
    "right_wrist",
    "left_pinky",
    "right_pinky",
    "left_index",
    "right_index",
    "left_thumb",
    "right_thumb",
    "left_hip",
    "right_hip",
    "left_knee",
    "right_knee",
    "left_ankle",
    "right_ankle",
    "left_heel",
    "right_heel",
    "left_foot_index",
    "right_foot_index",
]

KEYPOINT_FIELDS = [
    "bag",
    "segment_id",
    "label",
    "task_type",
    "NP3GAIT",
    "frame_index",
    "timestamp_s",
    "pose_detected",
    "landmark_index",
    "landmark",
    "pixel_x",
    "pixel_y",
    "visibility",
    "depth_m",
    "x_m",
    "y_m",
    "z_m",
    "valid_depth",
]

METRIC_NAMES = [
    "SP_U",
    "RA_AMP_U",
    "LA_AMP_U",
    "RA_STD_U",
    "LA_STD_U",
    "SYM_U",
    "R_JERK_U",
    "L_JERK_U",
    "ASA_U",
    "ASYM_IND_U",
    "TRA_U",
    "T_AMP_U",
    "STR_T_U",
    "STR_CV_U",
    "STEP_REG_U",
    "STEP_SYM_U",
    "JERK_T_U",
]

METRIC_FIELDS = [
    "bag",
    "segment_id",
    "label",
    "task_type",
    "NP3GAIT",
    "start_s",
    "end_s",
    "duration_s",
    "metric_status",
    "frames_total",
    "pose_detection_rate",
    "valid_depth_rate",
] + METRIC_NAMES

POSE_MODEL_BUNDLES = {
    0: (
        "pose_landmarker_lite.task",
        "https://storage.googleapis.com/mediapipe-models/pose_landmarker/pose_landmarker_lite/float16/latest/pose_landmarker_lite.task",
    ),
    1: (
        "pose_landmarker_full.task",
        "https://storage.googleapis.com/mediapipe-models/pose_landmarker/pose_landmarker_full/float16/latest/pose_landmarker_full.task",
    ),
    2: (
        "pose_landmarker_heavy.task",
        "https://storage.googleapis.com/mediapipe-models/pose_landmarker/pose_landmarker_heavy/float16/latest/pose_landmarker_heavy.task",
    ),
}


@dataclass(frozen=True)
class Segment:
    segment_id: str
    label: str
    task_type: str
    start_s: float
    end_s: float

    @property
    def duration_s(self) -> float:
        return max(0.0, self.end_s - self.start_s)

    @property
    def safe_name(self) -> str:
        return safe_filename(f"{self.segment_id}_{self.label}")


@dataclass
class LandmarkPoint:
    index: int
    name: str
    pixel_x: Optional[float] = None
    pixel_y: Optional[float] = None
    visibility: Optional[float] = None
    depth_m: Optional[float] = None
    x_m: Optional[float] = None
    y_m: Optional[float] = None
    z_m: Optional[float] = None
    valid_depth: bool = False

    def xyz(self) -> Optional[np.ndarray]:
        if not self.valid_depth:
            return None
        if self.x_m is None or self.y_m is None or self.z_m is None:
            return None
        return np.array([self.x_m, self.y_m, self.z_m], dtype=float)


@dataclass
class PoseFrame:
    frame_index: int
    timestamp_s: float
    detected: bool
    landmarks: Dict[str, LandmarkPoint] = field(default_factory=dict)


@dataclass
class SegmentQuality:
    bag: str
    segment_id: str
    label: str
    task_type: str
    start_s: float
    end_s: float
    frames_total: int
    frames_with_pose: int
    pose_detection_rate: Optional[float]
    valid_depth_points: int
    landmark_slots: int
    valid_depth_rate: Optional[float]
    keypoints_csv: str
    annotated_video: str


class SegmentOutput:
    def __init__(
        self,
        bag_name: str,
        segment: Segment,
        out_dir: Path,
        fps: float,
        frame_size: Tuple[int, int],
        codec: str,
        true_score: Optional[float] = None,
        write_annotated_video: bool = True,
    ):
        self.bag_name = bag_name
        self.segment = segment
        self.true_score = true_score
        self.fps = fps if fps > 0 else 30.0
        self.frame_size = frame_size
        self.codec = codec
        self.write_annotated_video = write_annotated_video
        output_stem = f"{safe_filename(bag_name)}_{segment.safe_name}"
        self.keypoints_path = out_dir / "keypoints" / f"{output_stem}_keypoints.csv"
        video_suffix = ".webm" if codec.upper() in {"VP80", "VP90"} else ".mp4"
        self.video_path = out_dir / "videos" / f"{output_stem}_annotated{video_suffix}"
        self.keypoints_file = self.keypoints_path.open("w", newline="", encoding="utf-8-sig")
        self.keypoints_writer = csv.DictWriter(self.keypoints_file, fieldnames=KEYPOINT_FIELDS)
        self.keypoints_writer.writeheader()
        self.video_writer = None
        self.frames: List[PoseFrame] = []
        self.frames_total = 0
        self.frames_with_pose = 0
        self.valid_depth_points = 0
        self.landmark_slots = 0

    def write_frame(self, frame: PoseFrame, annotated_frame) -> None:
        self.frames_total += 1
        self.landmark_slots += len(LANDMARK_NAMES)
        if frame.detected:
            self.frames_with_pose += 1
        self.valid_depth_points += sum(1 for point in frame.landmarks.values() if point.valid_depth)
        self.frames.append(frame)

        for index, name in enumerate(LANDMARK_NAMES):
            point = frame.landmarks.get(name)
            self.keypoints_writer.writerow(
                {
                    "bag": self.bag_name,
                    "segment_id": self.segment.segment_id,
                    "label": self.segment.label,
                    "task_type": self.segment.task_type,
                    "NP3GAIT": clean_cell(self.true_score),
                    "frame_index": frame.frame_index,
                    "timestamp_s": format_float(frame.timestamp_s, 6),
                    "pose_detected": int(frame.detected),
                    "landmark_index": index,
                    "landmark": name,
                    "pixel_x": format_float(point.pixel_x if point else None, 3),
                    "pixel_y": format_float(point.pixel_y if point else None, 3),
                    "visibility": format_float(point.visibility if point else None, 6),
                    "depth_m": format_float(point.depth_m if point else None, 6),
                    "x_m": format_float(point.x_m if point else None, 6),
                    "y_m": format_float(point.y_m if point else None, 6),
                    "z_m": format_float(point.z_m if point else None, 6),
                    "valid_depth": int(bool(point and point.valid_depth)),
                }
            )

        if not self.write_annotated_video or annotated_frame is None:
            return
        if self.video_writer is None:
            self.video_path.parent.mkdir(parents=True, exist_ok=True)
            import cv2

            fourcc = cv2.VideoWriter_fourcc(*self.codec)
            self.video_writer = cv2.VideoWriter(str(self.video_path), fourcc, self.fps, self.frame_size)
            if not self.video_writer.isOpened():
                raise RuntimeError(f"Could not open video writer for {self.video_path}")
        self.video_writer.write(annotated_frame)

    def close(self) -> None:
        self.keypoints_file.close()
        if self.video_writer is not None:
            self.video_writer.release()

    def quality(self) -> SegmentQuality:
        pose_rate = safe_divide(self.frames_with_pose, self.frames_total)
        valid_rate = safe_divide(self.valid_depth_points, self.landmark_slots)
        return SegmentQuality(
            bag=self.bag_name,
            segment_id=self.segment.segment_id,
            label=self.segment.label,
            task_type=self.segment.task_type,
            start_s=self.segment.start_s,
            end_s=self.segment.end_s,
            frames_total=self.frames_total,
            frames_with_pose=self.frames_with_pose,
            pose_detection_rate=pose_rate,
            valid_depth_points=self.valid_depth_points,
            landmark_slots=self.landmark_slots,
            valid_depth_rate=valid_rate,
            keypoints_csv=str(self.keypoints_path),
            annotated_video=str(self.video_path if self.video_writer is not None else ""),
        )


def safe_filename(value: str) -> str:
    value = value.strip()
    value = re.sub(r"[^\w.-]+", "_", value, flags=re.UNICODE)
    return value.strip("._") or "segment"


def format_float(value: Optional[float], digits: int = 6) -> str:
    if value is None:
        return ""
    try:
        numeric = float(value)
    except (TypeError, ValueError):
        return ""
    if not math.isfinite(numeric):
        return ""
    return f"{numeric:.{digits}f}"


def safe_divide(numerator: float, denominator: float) -> Optional[float]:
    if denominator is None or denominator == 0:
        return None
    result = numerator / denominator
    return result if math.isfinite(result) else None


def clean_cell(value):
    if isinstance(value, float):
        if not math.isfinite(value):
            return ""
        return value
    if value is None:
        return ""
    return value


def parse_segments(path: Path) -> List[Segment]:
    required = {"segment_id", "label", "task_type", "start_s", "end_s"}
    with path.open("r", newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        if reader.fieldnames is None:
            raise ValueError(f"Segment file has no header: {path}")
        missing = required - set(reader.fieldnames)
        if missing:
            raise ValueError(f"Segment file is missing columns: {', '.join(sorted(missing))}")
        segments = []
        for row_number, row in enumerate(reader, start=2):
            try:
                start_s = float(row["start_s"])
                end_s = float(row["end_s"])
            except ValueError as exc:
                raise ValueError(f"Invalid start_s/end_s at row {row_number}") from exc
            if start_s < 0 or end_s <= start_s:
                raise ValueError(f"Invalid time range at row {row_number}: start_s={start_s}, end_s={end_s}")
            segment_id = str(row["segment_id"]).strip()
            label = str(row["label"]).strip()
            task_type = str(row["task_type"]).strip().lower()
            if not segment_id or not label or not task_type:
                raise ValueError(f"segment_id, label, and task_type are required at row {row_number}")
            segments.append(Segment(segment_id=segment_id, label=label, task_type=task_type, start_s=start_s, end_s=end_s))
    if not segments:
        raise ValueError("Segment file contains no segment rows")
    return sorted(segments, key=lambda item: item.start_s)


def normalize_subject_name(value: object) -> str:
    text = "" if value is None else str(value)
    text = re.sub(r"\s+", "", text.strip())
    return text


def normalize_header(value: object) -> str:
    return normalize_subject_name(value).lower().replace("_", "")


def coerce_score(value: object) -> Optional[float]:
    if value is None or value == "":
        return None
    try:
        numeric = float(value)
    except (TypeError, ValueError):
        return None
    if not math.isfinite(numeric):
        return None
    return int(numeric) if numeric.is_integer() else numeric


def score_rows_to_map(rows: Sequence[Sequence[object]], source: Path) -> Dict[str, float]:
    if not rows:
        return {}
    headers = [normalize_header(value) for value in rows[0]]
    name_aliases = {"name", "姓名", "患者", "病人", "patient", "patientname", "subject"}
    score_aliases = {"score", "评分", "真实标签", "标签", "np3gait", "true", "label"}
    try:
        name_index = next(index for index, header in enumerate(headers) if header in name_aliases)
        score_index = next(index for index, header in enumerate(headers) if header in score_aliases)
    except StopIteration as exc:
        raise ValueError(f"Score file must contain name and score columns: {source}") from exc

    score_map: Dict[str, float] = {}
    for row in rows[1:]:
        if name_index >= len(row) or score_index >= len(row):
            continue
        name = normalize_subject_name(row[name_index])
        score = coerce_score(row[score_index])
        if name and score is not None:
            score_map[name] = score
    return score_map


def load_score_map(path: Path) -> Dict[str, float]:
    data = path.read_bytes()
    if data.startswith(b"PK"):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("openpyxl is required to read Excel score files.") from exc
        workbook = load_workbook(BytesIO(data), data_only=True, read_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        return score_rows_to_map(rows, path)

    for encoding in ("utf-8-sig", "gbk"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            text = ""
    if not text:
        raise ValueError(f"Could not decode score file: {path}")
    reader = csv.reader(text.splitlines())
    return score_rows_to_map(list(reader), path)


def resolve_score_path(value: Optional[str]) -> Optional[Path]:
    if not value:
        return None
    raw_path = Path(value).expanduser()
    candidates = [raw_path]
    if not raw_path.is_absolute():
        script_dir = Path(__file__).resolve().parent
        candidates = [Path.cwd() / raw_path, script_dir / raw_path]
    for candidate in candidates:
        if candidate.exists():
            return candidate.resolve()
    return None


def score_for_bag(bag_name: str, score_map: Dict[str, float]) -> Optional[float]:
    normalized_bag = normalize_subject_name(bag_name)
    if normalized_bag in score_map:
        return score_map[normalized_bag]
    matches = [score for name, score in score_map.items() if name and name in normalized_bag]
    if len(matches) == 1:
        return matches[0]
    return None


def active_segments_at(segments: Sequence[Segment], timestamp_s: float) -> List[Segment]:
    return [segment for segment in segments if segment.start_s <= timestamp_s <= segment.end_s]


def ensure_output_dirs(out_dir: Path) -> None:
    for child in ["keypoints", "metrics", "reports", "videos"]:
        (out_dir / child).mkdir(parents=True, exist_ok=True)


def resolve_model_path(value: Optional[str], model_complexity: int) -> Path:
    if value:
        return Path(value).expanduser().resolve()
    model_name, _ = POSE_MODEL_BUNDLES[model_complexity]
    return (Path(__file__).resolve().parent / "models" / model_name).resolve()


def ensure_pose_model(model_path: Path, model_complexity: int) -> Path:
    if model_path.exists():
        return model_path
    _, model_url = POSE_MODEL_BUNDLES[model_complexity]
    model_path.parent.mkdir(parents=True, exist_ok=True)
    temp_path = model_path.with_suffix(model_path.suffix + ".tmp")
    print(f"Downloading MediaPipe Pose Landmarker model: {model_url}")
    urllib.request.urlretrieve(model_url, temp_path)
    temp_path.replace(model_path)
    return model_path


def import_runtime_modules():
    try:
        import cv2
        import mediapipe as mp
        from mediapipe.tasks import python as mp_python
        from mediapipe.tasks.python import vision
        import pyrealsense2 as rs
    except ImportError as exc:
        missing = getattr(exc, "name", str(exc))
        raise RuntimeError(
            "Missing runtime dependency. On Windows install requirements-win.txt "
            f"before processing bag files. Import failed for: {missing}"
        ) from exc
    return cv2, mp, vision, mp_python.BaseOptions, rs


def get_video_intrinsics(video_profile):
    if hasattr(video_profile, "get_intrinsics"):
        return video_profile.get_intrinsics()
    if hasattr(video_profile, "intrinsics"):
        return video_profile.intrinsics
    raise RuntimeError("Could not read color stream intrinsics")


def color_frame_to_bgr(color_frame, cv2, np_module, rs):
    data = np_module.asanyarray(color_frame.get_data())
    fmt = color_frame.get_profile().as_video_stream_profile().format()
    if fmt == rs.format.bgr8:
        return data.copy()
    if fmt == rs.format.rgb8:
        return cv2.cvtColor(data, cv2.COLOR_RGB2BGR)
    if fmt == rs.format.rgba8:
        return cv2.cvtColor(data, cv2.COLOR_RGBA2BGR)
    if fmt == rs.format.bgra8:
        return cv2.cvtColor(data, cv2.COLOR_BGRA2BGR)
    if fmt == rs.format.yuyv:
        return cv2.cvtColor(data, cv2.COLOR_YUV2BGR_YUY2)
    if data.ndim == 3 and data.shape[2] == 3:
        return data.copy()
    raise RuntimeError(f"Unsupported color frame format: {fmt}")


def depth_at_pixel(depth_frame, pixel_x: float, pixel_y: float, width: int, height: int, window: int) -> Tuple[Optional[float], bool]:
    if not math.isfinite(pixel_x) or not math.isfinite(pixel_y):
        return None, False
    x = int(round(pixel_x))
    y = int(round(pixel_y))
    if x < 0 or y < 0 or x >= width or y >= height:
        return None, False

    depth = float(depth_frame.get_distance(x, y))
    if depth > 0 and math.isfinite(depth):
        return depth, True

    radius = max(0, int(window) // 2)
    if radius <= 0:
        return None, False

    values = []
    for yy in range(max(0, y - radius), min(height, y + radius + 1)):
        for xx in range(max(0, x - radius), min(width, x + radius + 1)):
            candidate = float(depth_frame.get_distance(xx, yy))
            if candidate > 0 and math.isfinite(candidate):
                values.append(candidate)
    if not values:
        return None, False
    return float(statistics.median(values)), True


def build_pose_frame(results, depth_frame, intrinsics, frame_index: int, timestamp_s: float, width: int, height: int, depth_window: int, rs) -> PoseFrame:
    if not results.pose_landmarks:
        return PoseFrame(frame_index=frame_index, timestamp_s=timestamp_s, detected=False)

    points: Dict[str, LandmarkPoint] = {}
    for index, landmark in enumerate(results.pose_landmarks[0]):
        name = LANDMARK_NAMES[index]
        pixel_x = float(landmark.x * width)
        pixel_y = float(landmark.y * height)
        depth_m, valid = depth_at_pixel(depth_frame, pixel_x, pixel_y, width, height, depth_window)
        x_m = y_m = z_m = None
        if valid and depth_m is not None:
            try:
                point_3d = rs.rs2_deproject_pixel_to_point(intrinsics, [pixel_x, pixel_y], depth_m)
                x_m, y_m, z_m = map(float, point_3d)
            except Exception:
                valid = False
                depth_m = None
        points[name] = LandmarkPoint(
            index=index,
            name=name,
            pixel_x=pixel_x,
            pixel_y=pixel_y,
            visibility=float(landmark.visibility) if landmark.visibility is not None else None,
            depth_m=depth_m,
            x_m=x_m,
            y_m=y_m,
            z_m=z_m,
            valid_depth=bool(valid),
        )
    return PoseFrame(frame_index=frame_index, timestamp_s=timestamp_s, detected=True, landmarks=points)


def annotate_frame(frame_bgr, results, vision, cv2, segment: Segment, frame: PoseFrame, valid_rate: Optional[float]) -> np.ndarray:
    annotated = frame_bgr.copy()
    if results.pose_landmarks:
        landmarks = results.pose_landmarks[0]
        height, width = annotated.shape[:2]
        # MediaPipe Tasks does not expose the legacy drawing_utils module.
        # Draw directly from its official connection list so this remains
        # compatible with the version used by all integrated modules.
        for connection in vision.PoseLandmarksConnections.POSE_LANDMARKS:
            start = landmarks[connection.start]
            end = landmarks[connection.end]
            start_visibility = float(start.visibility or 0.0)
            end_visibility = float(end.visibility or 0.0)
            if min(start_visibility, end_visibility) < 0.35:
                continue
            start_point = (int(start.x * width), int(start.y * height))
            end_point = (int(end.x * width), int(end.y * height))
            cv2.line(
                annotated,
                start_point,
                end_point,
                (70, 220, 205),
                3,
                cv2.LINE_AA,
            )
        for landmark in landmarks:
            if float(landmark.visibility or 0.0) < 0.35:
                continue
            point = (int(landmark.x * width), int(landmark.y * height))
            cv2.circle(annotated, point, 4, (255, 255, 255), -1, cv2.LINE_AA)
            cv2.circle(annotated, point, 5, (14, 139, 130), 1, cv2.LINE_AA)

    # 视频只绘制姿态骨架；片段、时间、检出率和深度质量改由前端
    # 播放器侧栏展示，避免逐帧文字遮挡动作画面。
    return annotated


def get_point(frame: PoseFrame, name: str) -> Optional[np.ndarray]:
    point = frame.landmarks.get(name)
    if point is None:
        return None
    return point.xyz()


def normalize(vector: np.ndarray) -> Optional[np.ndarray]:
    norm = float(np.linalg.norm(vector))
    if norm <= 1e-9 or not math.isfinite(norm):
        return None
    return vector / norm


def body_axes(frame: PoseFrame) -> Optional[Tuple[np.ndarray, np.ndarray, np.ndarray]]:
    left_hip = get_point(frame, "left_hip")
    right_hip = get_point(frame, "right_hip")
    left_shoulder = get_point(frame, "left_shoulder")
    right_shoulder = get_point(frame, "right_shoulder")
    if left_hip is None or right_hip is None or left_shoulder is None or right_shoulder is None:
        return None

    hip_mid = (left_hip + right_hip) * 0.5
    shoulder_mid = (left_shoulder + right_shoulder) * 0.5
    lateral = normalize(right_hip - left_hip)
    if lateral is None:
        lateral = normalize(right_shoulder - left_shoulder)
    if lateral is None:
        return None

    vertical = shoulder_mid - hip_mid
    vertical = vertical - np.dot(vertical, lateral) * lateral
    vertical = normalize(vertical)
    if vertical is None:
        return None

    forward = normalize(np.cross(lateral, vertical))
    if forward is None:
        return None
    return lateral, vertical, forward


def midpoint(frame: PoseFrame, first: str, second: str) -> Optional[np.ndarray]:
    a = get_point(frame, first)
    b = get_point(frame, second)
    if a is None or b is None:
        return None
    return (a + b) * 0.5


def trunk_center(frame: PoseFrame) -> Optional[np.ndarray]:
    shoulder = midpoint(frame, "left_shoulder", "right_shoulder")
    hip = midpoint(frame, "left_hip", "right_hip")
    if shoulder is None and hip is None:
        return None
    if shoulder is None:
        return hip
    if hip is None:
        return shoulder
    return (shoulder + hip) * 0.5


def median_fps(frames: Sequence[PoseFrame]) -> float:
    times = np.array([frame.timestamp_s for frame in frames], dtype=float)
    if len(times) < 3:
        return 30.0
    diffs = np.diff(times)
    diffs = diffs[diffs > 1e-6]
    if len(diffs) == 0:
        return 30.0
    fps = 1.0 / float(np.median(diffs))
    if not math.isfinite(fps) or fps <= 0:
        return 30.0
    return fps


def moving_average(values: Sequence[float], window: int = 5) -> np.ndarray:
    array = np.asarray(values, dtype=float)
    if len(array) == 0:
        return array
    window = int(max(1, min(window, len(array))))
    if window == 1:
        return array.copy()
    kernel = np.ones(window, dtype=float) / window
    pad_left = window // 2
    pad_right = window - 1 - pad_left
    padded = np.pad(array, (pad_left, pad_right), mode="edge")
    return np.convolve(padded, kernel, mode="valid")


def local_extrema(values: Sequence[float], min_distance: int, mode: str) -> List[int]:
    array = np.asarray(values, dtype=float)
    if len(array) < 3:
        return []
    candidates = []
    for index in range(1, len(array) - 1):
        if mode == "max" and array[index] >= array[index - 1] and array[index] > array[index + 1]:
            candidates.append(index)
        elif mode == "min" and array[index] <= array[index - 1] and array[index] < array[index + 1]:
            candidates.append(index)
    if not candidates:
        return []

    selected: List[int] = []
    for candidate in candidates:
        if not selected:
            selected.append(candidate)
            continue
        if candidate - selected[-1] >= min_distance:
            selected.append(candidate)
            continue
        if mode == "max" and array[candidate] > array[selected[-1]]:
            selected[-1] = candidate
        elif mode == "min" and array[candidate] < array[selected[-1]]:
            selected[-1] = candidate
    return selected


def swing_amplitudes(values: Sequence[float], fps: float) -> List[float]:
    if len(values) < 5:
        return []
    smooth = moving_average(values, window=max(3, int(round(fps * 0.15))))
    min_distance = max(3, int(round(fps * 0.25)))
    extrema = [(index, "max") for index in local_extrema(smooth, min_distance, "max")]
    extrema += [(index, "min") for index in local_extrema(smooth, min_distance, "min")]
    extrema.sort()
    amplitudes = []
    for (left_index, left_type), (right_index, right_type) in zip(extrema, extrema[1:]):
        if left_type == right_type:
            continue
        amplitude = abs(float(smooth[right_index] - smooth[left_index]))
        if amplitude > 1e-6:
            amplitudes.append(amplitude)

    if amplitudes:
        global_range = float(np.nanpercentile(smooth, 95) - np.nanpercentile(smooth, 5))
        threshold = max(1.0, 0.05 * global_range)
        amplitudes = [item for item in amplitudes if item >= threshold]
    if amplitudes:
        return amplitudes

    fallback = float(np.nanpercentile(smooth, 95) - np.nanpercentile(smooth, 5))
    return [fallback] if fallback > 0 else []


def arm_angle_series(frames: Sequence[PoseFrame], side: str) -> Tuple[List[float], List[float]]:
    shoulder_name = f"{side}_shoulder"
    wrist_name = f"{side}_wrist"
    times: List[float] = []
    angles: List[float] = []
    for frame in frames:
        axes = body_axes(frame)
        shoulder = get_point(frame, shoulder_name)
        wrist = get_point(frame, wrist_name)
        if axes is None or shoulder is None or wrist is None:
            continue
        _, vertical, forward = axes
        arm = wrist - shoulder
        anterior = float(np.dot(arm, forward))
        vertical_component = float(np.dot(arm, vertical))
        angle = math.degrees(math.atan2(anterior, abs(vertical_component) + 1e-9))
        if math.isfinite(angle):
            times.append(frame.timestamp_s)
            angles.append(angle)
    return times, angles


def scalar_jerk_rms(times: Sequence[float], values: Sequence[float]) -> Optional[float]:
    if len(times) < 6 or len(values) < 6:
        return None
    time_array = np.asarray(times, dtype=float)
    value_array = moving_average(values, window=5)
    order = np.argsort(time_array)
    time_array = time_array[order]
    value_array = value_array[order]
    unique = np.concatenate(([True], np.diff(time_array) > 1e-6))
    time_array = time_array[unique]
    value_array = value_array[unique]
    if len(time_array) < 6:
        return None
    try:
        first = np.gradient(value_array, time_array)
        second = np.gradient(first, time_array)
        third = np.gradient(second, time_array)
    except Exception:
        return None
    rms = float(np.sqrt(np.nanmean(np.square(third))))
    return rms if math.isfinite(rms) else None


def vector_jerk_rms(times: Sequence[float], positions: Sequence[np.ndarray]) -> Optional[float]:
    if len(times) < 6 or len(positions) < 6:
        return None
    time_array = np.asarray(times, dtype=float)
    position_array = np.asarray(positions, dtype=float)
    order = np.argsort(time_array)
    time_array = time_array[order]
    position_array = position_array[order]
    unique = np.concatenate(([True], np.diff(time_array) > 1e-6))
    time_array = time_array[unique]
    position_array = position_array[unique]
    if len(time_array) < 6:
        return None
    try:
        for axis in range(position_array.shape[1]):
            position_array[:, axis] = moving_average(position_array[:, axis], window=5)
        first = np.gradient(position_array, time_array, axis=0)
        second = np.gradient(first, time_array, axis=0)
        third = np.gradient(second, time_array, axis=0)
    except Exception:
        return None
    magnitude = np.linalg.norm(third, axis=1)
    rms = float(np.sqrt(np.nanmean(np.square(magnitude))))
    return rms if math.isfinite(rms) else None


def arm_metrics(frames: Sequence[PoseFrame], side: str) -> Tuple[Optional[float], Optional[float], Optional[float]]:
    fps = median_fps(frames)
    times, angles = arm_angle_series(frames, side)
    if len(angles) < 5:
        return None, None, None
    amplitudes = swing_amplitudes(angles, fps)
    if not amplitudes:
        return None, None, scalar_jerk_rms(times, angles)
    amp_mean = float(np.mean(amplitudes))
    amp_std = float(np.std(amplitudes, ddof=1)) if len(amplitudes) > 1 else 0.0
    jerk = scalar_jerk_rms(times, angles)
    return amp_mean, amp_std, jerk


def trunk_yaw_series(frames: Sequence[PoseFrame]) -> Tuple[List[float], List[float]]:
    times: List[float] = []
    yaw_values: List[float] = []
    for frame in frames:
        left_shoulder = get_point(frame, "left_shoulder")
        right_shoulder = get_point(frame, "right_shoulder")
        left_hip = get_point(frame, "left_hip")
        right_hip = get_point(frame, "right_hip")
        if left_shoulder is None or right_shoulder is None or left_hip is None or right_hip is None:
            continue
        shoulder_axis = right_shoulder - left_shoulder
        hip_axis = right_hip - left_hip
        lateral = shoulder_axis + hip_axis
        if np.linalg.norm(lateral[[0, 2]]) <= 1e-9:
            continue
        yaw = math.degrees(math.atan2(float(lateral[2]), float(lateral[0])))
        if math.isfinite(yaw):
            times.append(frame.timestamp_s)
            yaw_values.append(yaw)
    if len(yaw_values) >= 2:
        yaw_values = list(np.degrees(np.unwrap(np.radians(yaw_values))))
    return times, yaw_values


def asymmetry_score(first: Optional[float], second: Optional[float]) -> Optional[float]:
    if first is None or second is None or first <= 0 or second <= 0:
        return None
    greater = max(first, second)
    smaller = min(first, second)
    angle = abs(45.0 - math.degrees(math.atan(greater / smaller)))
    return angle * 100.0 / 90.0


def trunk_metrics(frames: Sequence[PoseFrame]) -> Tuple[Optional[float], Optional[float]]:
    _, yaw_values = trunk_yaw_series(frames)
    if len(yaw_values) < 5:
        return None, None
    smooth = moving_average(yaw_values, window=5)
    centered = smooth - np.nanmedian(smooth)
    left_mag = abs(float(np.nanpercentile(centered, 5)))
    right_mag = abs(float(np.nanpercentile(centered, 95)))
    amplitude = float(np.nanpercentile(smooth, 95) - np.nanpercentile(smooth, 5))
    tra = asymmetry_score(left_mag, right_mag)
    return tra, amplitude if math.isfinite(amplitude) else None


def pelvis_center(frame: PoseFrame) -> Optional[np.ndarray]:
    return midpoint(frame, "left_hip", "right_hip")


def infer_walk_axis(frames: Sequence[PoseFrame]) -> np.ndarray:
    centers = [(frame.timestamp_s, pelvis_center(frame)) for frame in frames]
    centers = [(timestamp, center) for timestamp, center in centers if center is not None]
    if len(centers) >= 2:
        start = centers[0][1]
        end = centers[-1][1]
        displacement = end - start
        displacement[1] = 0.0
        axis = normalize(displacement)
        if axis is not None:
            return axis
    return np.array([0.0, 0.0, 1.0], dtype=float)


def foot_signal(frames: Sequence[PoseFrame], side: str, axis: np.ndarray) -> Tuple[List[float], List[float]]:
    times: List[float] = []
    values: List[float] = []
    foot_names = [f"{side}_heel", f"{side}_ankle", f"{side}_foot_index"]
    for frame in frames:
        pelvis = pelvis_center(frame)
        foot = None
        for name in foot_names:
            foot = get_point(frame, name)
            if foot is not None:
                break
        if pelvis is None or foot is None:
            continue
        relative = foot - pelvis
        values.append(float(np.dot(relative, axis)))
        times.append(frame.timestamp_s)
    return times, values


def event_times_from_signal(times: Sequence[float], values: Sequence[float], fps: float) -> List[float]:
    if len(values) < 5:
        return []
    smooth = moving_average(values, window=max(3, int(round(fps * 0.12))))
    min_distance = max(3, int(round(fps * 0.35)))
    maxima = local_extrema(smooth, min_distance, "max")
    minima = local_extrema(smooth, min_distance, "min")
    chosen = maxima if len(maxima) >= len(minima) else minima
    return [float(times[index]) for index in chosen]


def filtered_durations(values: Sequence[float], min_s: float, max_s: float) -> List[float]:
    return [float(value) for value in values if min_s <= float(value) <= max_s]


def normalized_autocorrelation(values: Sequence[float]) -> np.ndarray:
    array = np.asarray(values, dtype=float)
    if len(array) < 3:
        return np.array([], dtype=float)
    array = array - np.nanmean(array)
    denom = float(np.dot(array, array))
    if denom <= 1e-12:
        return np.array([], dtype=float)
    corr = np.correlate(array, array, mode="full")[len(array) - 1 :]
    return corr / denom


def autocorr_peak_near(corr: np.ndarray, target_lag: int, window: int) -> Optional[float]:
    if len(corr) == 0 or target_lag <= 0:
        return None
    left = max(1, target_lag - window)
    right = min(len(corr), target_lag + window + 1)
    if right <= left:
        return None
    value = float(np.nanmax(corr[left:right]))
    return value if math.isfinite(value) else None


def vertical_acceleration_series(frames: Sequence[PoseFrame]) -> Tuple[List[float], List[float]]:
    times = []
    values = []
    for frame in frames:
        center = trunk_center(frame)
        if center is None:
            continue
        times.append(frame.timestamp_s)
        values.append(float(center[1]))
    if len(values) < 6:
        return [], []
    time_array = np.asarray(times, dtype=float)
    value_array = moving_average(values, window=5)
    unique = np.concatenate(([True], np.diff(time_array) > 1e-6))
    time_array = time_array[unique]
    value_array = value_array[unique]
    if len(value_array) < 6:
        return [], []
    try:
        velocity = np.gradient(value_array, time_array)
        acceleration = np.gradient(velocity, time_array)
    except Exception:
        return [], []
    return list(time_array), list(acceleration)


def step_regularities(frames: Sequence[PoseFrame], median_step_s: Optional[float], median_stride_s: Optional[float]) -> Tuple[Optional[float], Optional[float]]:
    times, acceleration = vertical_acceleration_series(frames)
    if len(acceleration) < 8:
        return None, None
    fps = median_fps([PoseFrame(i, t, True) for i, t in enumerate(times)])
    corr = normalized_autocorrelation(acceleration)
    if len(corr) == 0:
        return None, None

    if median_step_s is None or median_step_s <= 0:
        peaks = local_extrema(corr, max(2, int(round(fps * 0.2))), "max")
        peaks = [peak for peak in peaks if int(round(fps * 0.2)) <= peak <= int(round(fps * 2.0))]
        step_reg = float(corr[peaks[0]]) if peaks else None
        stride_reg = float(corr[peaks[1]]) if len(peaks) > 1 else None
    else:
        step_lag = int(round(median_step_s * fps))
        stride_lag = int(round((median_stride_s or median_step_s * 2.0) * fps))
        step_reg = autocorr_peak_near(corr, step_lag, max(2, int(round(fps * 0.15))))
        stride_reg = autocorr_peak_near(corr, stride_lag, max(2, int(round(fps * 0.2))))

    step_reg = max(0.0, step_reg) if step_reg is not None else None
    stride_reg = max(0.0, stride_reg) if stride_reg is not None else None
    step_sym = safe_divide(step_reg, stride_reg) if step_reg is not None and stride_reg is not None else None
    return step_reg, step_sym


def stride_metrics(frames: Sequence[PoseFrame]) -> Tuple[Optional[float], Optional[float], Optional[float], Optional[float]]:
    fps = median_fps(frames)
    axis = infer_walk_axis(frames)
    left_times, left_signal = foot_signal(frames, "left", axis)
    right_times, right_signal = foot_signal(frames, "right", axis)
    left_events = event_times_from_signal(left_times, left_signal, fps)
    right_events = event_times_from_signal(right_times, right_signal, fps)

    events = sorted([(time, "left") for time in left_events] + [(time, "right") for time in right_events])
    step_durations = filtered_durations([events[i + 1][0] - events[i][0] for i in range(len(events) - 1)], 0.2, 2.0)
    stride_durations = []
    for side_events in [left_events, right_events]:
        stride_durations.extend(filtered_durations(np.diff(side_events), 0.5, 3.0))

    stride_time = float(np.mean(stride_durations)) if stride_durations else None
    stride_cv = None
    if stride_durations and np.mean(stride_durations) > 0:
        stride_cv = float(np.std(stride_durations, ddof=1) / np.mean(stride_durations) * 100.0) if len(stride_durations) > 1 else 0.0

    median_step = float(np.median(step_durations)) if step_durations else None
    median_stride = float(np.median(stride_durations)) if stride_durations else None
    step_reg, step_sym = step_regularities(frames, median_step, median_stride)
    return stride_time, stride_cv, step_reg, step_sym


def trunk_position_series(frames: Sequence[PoseFrame]) -> Tuple[List[float], List[np.ndarray]]:
    times: List[float] = []
    positions: List[np.ndarray] = []
    for frame in frames:
        center = trunk_center(frame)
        if center is None:
            continue
        times.append(frame.timestamp_s)
        positions.append(center)
    return times, positions


def compute_walk_metrics(segment: Segment, frames: Sequence[PoseFrame], walk_distance_m: float) -> Dict[str, Optional[float]]:
    metrics: Dict[str, Optional[float]] = {name: None for name in METRIC_NAMES}
    duration_s = segment.duration_s
    metrics["SP_U"] = safe_divide(walk_distance_m, duration_s)

    right_amp, right_std, right_jerk = arm_metrics(frames, "right")
    left_amp, left_std, left_jerk = arm_metrics(frames, "left")
    metrics["RA_AMP_U"] = right_amp
    metrics["LA_AMP_U"] = left_amp
    metrics["RA_STD_U"] = right_std
    metrics["LA_STD_U"] = left_std
    metrics["R_JERK_U"] = right_jerk
    metrics["L_JERK_U"] = left_jerk

    if right_amp is not None and left_amp is not None and left_amp > 0:
        metrics["SYM_U"] = 1.0 - (right_amp / left_amp)
    if right_amp is not None and left_amp is not None and (right_amp + left_amp) > 0:
        metrics["ASYM_IND_U"] = abs(right_amp - left_amp) / (right_amp + left_amp) * 100.0
    metrics["ASA_U"] = asymmetry_score(right_amp, left_amp)

    tra, trunk_amp = trunk_metrics(frames)
    metrics["TRA_U"] = tra
    metrics["T_AMP_U"] = trunk_amp

    stride_time, stride_cv, step_reg, step_sym = stride_metrics(frames)
    metrics["STR_T_U"] = stride_time
    metrics["STR_CV_U"] = stride_cv
    metrics["STEP_REG_U"] = step_reg
    metrics["STEP_SYM_U"] = step_sym

    trunk_times, trunk_positions = trunk_position_series(frames)
    metrics["JERK_T_U"] = vector_jerk_rms(trunk_times, trunk_positions)
    return metrics


def quality_to_row(quality: SegmentQuality) -> Dict[str, object]:
    return {
        "bag": quality.bag,
        "segment_id": quality.segment_id,
        "label": quality.label,
        "task_type": quality.task_type,
        "start_s": quality.start_s,
        "end_s": quality.end_s,
        "frames_total": quality.frames_total,
        "frames_with_pose": quality.frames_with_pose,
        "pose_detection_rate": quality.pose_detection_rate,
        "valid_depth_points": quality.valid_depth_points,
        "landmark_slots": quality.landmark_slots,
        "valid_depth_rate": quality.valid_depth_rate,
        "keypoints_csv": quality.keypoints_csv,
        "annotated_video": quality.annotated_video,
    }


def build_metric_rows(
    bag_name: str,
    segments: Sequence[Segment],
    outputs: Dict[str, SegmentOutput],
    walk_distance_m: float,
    true_score: Optional[float] = None,
) -> List[Dict[str, object]]:
    rows = []
    for segment in segments:
        output = outputs[segment.segment_id]
        quality = output.quality()
        row: Dict[str, object] = {
            "bag": bag_name,
            "segment_id": segment.segment_id,
            "label": segment.label,
            "task_type": segment.task_type,
            "NP3GAIT": true_score,
            "start_s": segment.start_s,
            "end_s": segment.end_s,
            "duration_s": segment.duration_s,
            "metric_status": "computed" if segment.task_type == "walk" else "not_applicable",
            "frames_total": quality.frames_total,
            "pose_detection_rate": quality.pose_detection_rate,
            "valid_depth_rate": quality.valid_depth_rate,
        }
        row.update({name: None for name in METRIC_NAMES})
        if segment.task_type == "walk":
            row.update(compute_walk_metrics(segment, output.frames, walk_distance_m))
        rows.append(row)
    return rows


def write_metrics_csv(path: Path, rows: Sequence[Dict[str, object]]) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=METRIC_FIELDS)
        writer.writeheader()
        for row in rows:
            writer.writerow({field: clean_cell(row.get(field)) for field in METRIC_FIELDS})


def keypoint_summary_rows(segments: Sequence[Segment], outputs: Dict[str, SegmentOutput]) -> List[Dict[str, object]]:
    rows: List[Dict[str, object]] = []
    for segment in segments:
        frames = outputs[segment.segment_id].frames
        for landmark_name in LANDMARK_NAMES:
            positions = []
            visible_values = []
            for frame in frames:
                point = frame.landmarks.get(landmark_name)
                if point is None:
                    continue
                if point.visibility is not None:
                    visible_values.append(point.visibility)
                xyz = point.xyz()
                if xyz is not None:
                    positions.append(xyz)
            valid_count = len(positions)
            mean_xyz = np.mean(positions, axis=0) if positions else [None, None, None]
            rows.append(
                {
                    "segment_id": segment.segment_id,
                    "label": segment.label,
                    "task_type": segment.task_type,
                    "landmark": landmark_name,
                    "frames_total": len(frames),
                    "valid_3d_frames": valid_count,
                    "valid_3d_rate": safe_divide(valid_count, len(frames)),
                    "mean_visibility": float(np.mean(visible_values)) if visible_values else None,
                    "mean_x_m": mean_xyz[0],
                    "mean_y_m": mean_xyz[1],
                    "mean_z_m": mean_xyz[2],
                }
            )
    return rows


def write_sheet(sheet, rows: Sequence[Dict[str, object]], fields: Sequence[str]) -> None:
    sheet.append(list(fields))
    for row in rows:
        sheet.append([clean_cell(row.get(field)) for field in fields])
    sheet.freeze_panes = "A2"


def autosize_sheet(sheet) -> None:
    for column_cells in sheet.columns:
        max_len = 0
        column_letter = column_cells[0].column_letter
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(value))
        sheet.column_dimensions[column_letter].width = min(max_len + 2, 60)


def write_excel_report(
    path: Path,
    segments: Sequence[Segment],
    metric_rows: Sequence[Dict[str, object]],
    quality_rows: Sequence[Dict[str, object]],
    summary_rows: Sequence[Dict[str, object]],
    true_score: Optional[float] = None,
) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError as exc:
        raise RuntimeError("openpyxl is required for Excel export. Install requirements-win.txt on Windows.") from exc

    workbook = Workbook()
    segments_sheet = workbook.active
    segments_sheet.title = "Segments"
    segment_rows = [
        {
                "segment_id": segment.segment_id,
                "label": segment.label,
                "task_type": segment.task_type,
                "NP3GAIT": true_score,
                "start_s": segment.start_s,
                "end_s": segment.end_s,
                "duration_s": segment.duration_s,
        }
        for segment in segments
    ]
    write_sheet(segments_sheet, segment_rows, ["segment_id", "label", "task_type", "NP3GAIT", "start_s", "end_s", "duration_s"])

    metrics_sheet = workbook.create_sheet("Metrics")
    write_sheet(metrics_sheet, metric_rows, METRIC_FIELDS)

    quality_sheet = workbook.create_sheet("Quality")
    quality_fields = [
        "bag",
        "segment_id",
        "label",
        "task_type",
        "start_s",
        "end_s",
        "frames_total",
        "frames_with_pose",
        "pose_detection_rate",
        "valid_depth_points",
        "landmark_slots",
        "valid_depth_rate",
        "keypoints_csv",
        "annotated_video",
    ]
    write_sheet(quality_sheet, quality_rows, quality_fields)

    summary_sheet = workbook.create_sheet("KeypointSummary")
    summary_fields = [
        "segment_id",
        "label",
        "task_type",
        "landmark",
        "frames_total",
        "valid_3d_frames",
        "valid_3d_rate",
        "mean_visibility",
        "mean_x_m",
        "mean_y_m",
        "mean_z_m",
    ]
    write_sheet(summary_sheet, summary_rows, summary_fields)

    notes_sheet = workbook.create_sheet("MethodNotes")
    notes_sheet.append(["metric", "method_note"])
    notes = {
        "SP_U": "walk_distance_m / segment duration.",
        "RA_AMP_U / LA_AMP_U": "Mean swing amplitude from shoulder-to-wrist 3D angle in the torso sagittal frame.",
        "RA_STD_U / LA_STD_U": "Standard deviation of detected arm swing amplitudes.",
        "SYM_U": "1 - right arm amplitude / left arm amplitude.",
        "ASA_U": "Absolute 45-degree arctangent asymmetry score between arm amplitudes.",
        "ASYM_IND_U": "Absolute arm amplitude difference divided by summed arm amplitudes, percent.",
        "TRA_U": "Absolute 45-degree arctangent asymmetry score between left/right trunk yaw magnitudes.",
        "T_AMP_U": "95th to 5th percentile range of trunk yaw angle.",
        "STR_T_U / STR_CV_U": "Stride timing estimated from alternating foot anterior-posterior extrema.",
        "STEP_REG_U / STEP_SYM_U": "Normalized autocorrelation of vertical trunk acceleration near estimated step/stride lags.",
        "R_JERK_U / L_JERK_U": "RMS third derivative of arm swing angle.",
        "JERK_T_U": "RMS third derivative of trunk center 3D position.",
    }
    for metric, note in notes.items():
        notes_sheet.append([metric, note])
    notes_sheet.freeze_panes = "A2"

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for sheet in workbook.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
        autosize_sheet(sheet)

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def run_analysis(args: argparse.Namespace) -> Dict[str, object]:
    bag_path = Path(args.bag).expanduser().resolve()
    segments_path = Path(args.segments).expanduser().resolve()
    out_dir = Path(args.out).expanduser().resolve()
    if not bag_path.exists():
        raise FileNotFoundError(f"Bag file not found: {bag_path}")
    if not segments_path.exists():
        raise FileNotFoundError(f"Segments file not found: {segments_path}")

    segments = parse_segments(segments_path)
    ensure_output_dirs(out_dir)
    bag_name = bag_path.stem
    score_path = resolve_score_path(args.score_file)
    score_map = load_score_map(score_path) if score_path is not None else {}
    true_score = score_for_bag(bag_name, score_map)
    if args.score_file and score_path is None:
        print(f"Warning: score file not found: {args.score_file}")
    elif score_path is not None:
        print(f"Loaded {len(score_map)} score labels from {score_path}")
        if true_score is None:
            print(f"Warning: no NP3GAIT score matched bag name: {bag_name}")
        else:
            print(f"Matched NP3GAIT={true_score} for {bag_name}")

    cv2, mp, vision, BaseOptions, rs = import_runtime_modules()
    model_path = ensure_pose_model(resolve_model_path(args.model_path, args.model_complexity), args.model_complexity)

    pipeline = rs.pipeline()
    config = rs.config()
    config.enable_device_from_file(str(bag_path), repeat_playback=False)
    config.enable_stream(rs.stream.color)
    config.enable_stream(rs.stream.depth)

    pose = None
    outputs: Dict[str, SegmentOutput] = {}
    frame_index = 0
    first_timestamp_ms = None
    last_pose_timestamp_ms = -1
    max_end_s = max(segment.end_s for segment in segments)
    crop_region = getattr(args, "crop_region", None)
    write_annotated_video = bool(
        getattr(args, "write_annotated_video", True)
    )
    write_excel = bool(getattr(args, "write_excel_report", True))
    progress_callback = getattr(args, "progress_callback", None)

    try:
        profile = pipeline.start(config)
        playback = profile.get_device().as_playback()
        playback.set_real_time(False)
        align = rs.align(rs.stream.color)

        color_profile = profile.get_stream(rs.stream.color).as_video_stream_profile()
        intrinsics = get_video_intrinsics(color_profile)
        width = int(intrinsics.width)
        height = int(intrinsics.height)
        fps = float(color_profile.fps() or 30.0)

        if crop_region is not None:
            crop_x, crop_y, crop_width, crop_height = crop_region
            crop_right = crop_x + crop_width
            crop_bottom = crop_y + crop_height
            if (
                crop_x < 0
                or crop_y < 0
                or crop_width <= 0
                or crop_height <= 0
                or crop_right > width
                or crop_bottom > height
            ):
                raise ValueError(
                    "crop_region 超出视频画面范围："
                    f"区域=({crop_x}, {crop_y}, {crop_width}, {crop_height})，"
                    f"画面={width}x{height}"
                )

        for segment in segments:
            outputs[segment.segment_id] = SegmentOutput(
                bag_name=bag_name,
                segment=segment,
                out_dir=out_dir,
                fps=fps,
                frame_size=(width, height),
                codec=args.video_codec,
                true_score=true_score,
                write_annotated_video=write_annotated_video,
            )

        pose_options = vision.PoseLandmarkerOptions(
            base_options=BaseOptions(model_asset_buffer=model_path.read_bytes()),
            running_mode=vision.RunningMode.VIDEO,
            num_poses=1,
            min_pose_detection_confidence=args.min_detection_confidence,
            min_pose_presence_confidence=args.min_detection_confidence,
            min_tracking_confidence=args.min_tracking_confidence,
        )
        pose = vision.PoseLandmarker.create_from_options(pose_options)

        print(f"Processing {bag_path}")
        print(f"MediaPipe model: {model_path}")
        print(f"Loaded {len(segments)} segments. Color stream: {width}x{height} @ {fps:g} FPS")

        while True:
            if args.max_frames is not None and frame_index >= args.max_frames:
                print(f"Stopped after --max-frames={args.max_frames}")
                break
            try:
                frames = pipeline.wait_for_frames()
            except RuntimeError:
                print("Reached end of bag playback")
                break

            aligned_frames = align.process(frames)
            color_frame = aligned_frames.get_color_frame()
            depth_frame = aligned_frames.get_depth_frame()
            if not color_frame or not depth_frame:
                continue

            timestamp_ms = float(color_frame.get_timestamp())
            if first_timestamp_ms is None:
                first_timestamp_ms = timestamp_ms
            timestamp_s = (timestamp_ms - first_timestamp_ms) / 1000.0

            if timestamp_s > max_end_s and frame_index > 0:
                break

            active = active_segments_at(segments, timestamp_s)
            if not active:
                frame_index += 1
                continue

            frame_bgr = color_frame_to_bgr(color_frame, cv2, np, rs)
            detection_frame = frame_bgr
            if crop_region is not None:
                # 保持原始画幅和相机内参不变，只屏蔽 ROI 外的像素；这样
                # MediaPipe 只会考虑用户框选区域内的人体，同时深度坐标和
                # 标注视频仍能使用原始画面坐标。
                detection_frame = np.zeros_like(frame_bgr)
                detection_frame[
                    crop_y:crop_bottom, crop_x:crop_right
                ] = frame_bgr[crop_y:crop_bottom, crop_x:crop_right]
            frame_rgb = cv2.cvtColor(detection_frame, cv2.COLOR_BGR2RGB)
            mp_image = mp.Image(image_format=mp.ImageFormat.SRGB, data=frame_rgb)
            pose_timestamp_ms = int(round(timestamp_s * 1000.0))
            if pose_timestamp_ms <= last_pose_timestamp_ms:
                pose_timestamp_ms = last_pose_timestamp_ms + 1
            last_pose_timestamp_ms = pose_timestamp_ms
            results = pose.detect_for_video(mp_image, pose_timestamp_ms)

            pose_frame = build_pose_frame(
                results=results,
                depth_frame=depth_frame,
                intrinsics=intrinsics,
                frame_index=frame_index,
                timestamp_s=timestamp_s,
                width=width,
                height=height,
                depth_window=args.depth_window,
                rs=rs,
            )
            valid_rate = safe_divide(sum(1 for point in pose_frame.landmarks.values() if point.valid_depth), len(LANDMARK_NAMES))

            for segment in active:
                annotated = (
                    annotate_frame(
                        frame_bgr,
                        results,
                        vision,
                        cv2,
                        segment,
                        pose_frame,
                        valid_rate,
                    )
                    if write_annotated_video
                    else None
                )
                outputs[segment.segment_id].write_frame(
                    pose_frame, annotated
                )

            if frame_index % max(1, int(fps * 5)) == 0:
                print(f"Processed frame {frame_index} at {timestamp_s:.2f}s")
                if progress_callback is not None and max_end_s > 0:
                    progress_callback(
                        min(0.9, max(0.05, timestamp_s / max_end_s * 0.9)),
                        f"已提取步态特征 {timestamp_s:.1f}/{max_end_s:.1f} 秒",
                    )
            frame_index += 1

    finally:
        try:
            pipeline.stop()
        except Exception:
            pass
        if pose is not None:
            pose.close()
        for output in outputs.values():
            output.close()

    metric_rows = build_metric_rows(bag_name, segments, outputs, args.walk_distance_m, true_score=true_score)
    quality_rows = [quality_to_row(outputs[segment.segment_id].quality()) for segment in segments]
    summary_rows = keypoint_summary_rows(segments, outputs)

    metrics_path = out_dir / "metrics" / f"{bag_name}_metrics.csv"
    report_path = out_dir / "reports" / f"{bag_name}_analysis.xlsx"
    write_metrics_csv(metrics_path, metric_rows)
    if write_excel:
        write_excel_report(
            report_path,
            segments,
            metric_rows,
            quality_rows,
            summary_rows,
            true_score=true_score,
        )

    print("Done.")
    print(f"Keypoint CSV directory: {out_dir / 'keypoints'}")
    print(f"Metrics CSV: {metrics_path}")
    if write_excel:
        print(f"Excel report: {report_path}")
    if write_annotated_video:
        print(f"Annotated videos: {out_dir / 'videos'}")
    return {
        "metrics_path": metrics_path,
        "report_path": report_path if write_excel else None,
        "keypoints_dir": out_dir / "keypoints",
        "videos_dir": (
            out_dir / "videos" if write_annotated_video else None
        ),
        "metric_rows": metric_rows,
        "quality_rows": quality_rows,
    }


def synthetic_walk_frames(duration_s: float = 8.0, fps: float = 30.0) -> List[PoseFrame]:
    frames: List[PoseFrame] = []
    count = int(duration_s * fps)
    for frame_index in range(count):
        t = frame_index / fps
        phase = 2.0 * math.pi * 1.8 * t
        pelvis = np.array([0.0, 1.0 + 0.015 * math.sin(phase * 2.0), 0.12 * t])
        shoulder = pelvis + np.array([0.0, -0.55, 0.0])
        left_hip = pelvis + np.array([-0.16, 0.0, 0.0])
        right_hip = pelvis + np.array([0.16, 0.0, 0.0])
        left_shoulder = shoulder + np.array([-0.22, 0.0, 0.0])
        right_shoulder = shoulder + np.array([0.22, 0.0, 0.0])
        left_ankle = pelvis + np.array([-0.10, 0.82, 0.22 * math.sin(phase)])
        right_ankle = pelvis + np.array([0.10, 0.82, 0.22 * math.sin(phase + math.pi)])
        left_wrist = left_shoulder + np.array([-0.03, 0.48, 0.24 * math.sin(phase + math.pi)])
        right_wrist = right_shoulder + np.array([0.03, 0.48, 0.22 * math.sin(phase)])
        points = {
            "left_hip": left_hip,
            "right_hip": right_hip,
            "left_shoulder": left_shoulder,
            "right_shoulder": right_shoulder,
            "left_ankle": left_ankle,
            "right_ankle": right_ankle,
            "left_heel": left_ankle + np.array([0.0, 0.02, -0.04]),
            "right_heel": right_ankle + np.array([0.0, 0.02, -0.04]),
            "left_foot_index": left_ankle + np.array([0.0, 0.02, 0.08]),
            "right_foot_index": right_ankle + np.array([0.0, 0.02, 0.08]),
            "left_wrist": left_wrist,
            "right_wrist": right_wrist,
        }
        landmarks = {}
        for index, name in enumerate(LANDMARK_NAMES):
            xyz = points.get(name, pelvis)
            landmarks[name] = LandmarkPoint(
                index=index,
                name=name,
                pixel_x=100.0,
                pixel_y=100.0,
                visibility=0.99,
                depth_m=float(xyz[2]),
                x_m=float(xyz[0]),
                y_m=float(xyz[1]),
                z_m=float(xyz[2]),
                valid_depth=True,
            )
        frames.append(PoseFrame(frame_index=frame_index, timestamp_s=t, detected=True, landmarks=landmarks))
    return frames


class InMemoryOutput:
    def __init__(self, bag_name: str, segment: Segment, frames: Sequence[PoseFrame], out_dir: Path):
        self.bag_name = bag_name
        self.segment = segment
        self.frames = list(frames)
        output_stem = f"{safe_filename(bag_name)}_{segment.safe_name}"
        self.keypoints_path = out_dir / "keypoints" / f"{output_stem}_keypoints.csv"
        self.video_path = out_dir / "videos" / f"{output_stem}_annotated.mp4"

    def quality(self) -> SegmentQuality:
        frames_total = len(self.frames)
        frames_with_pose = sum(1 for frame in self.frames if frame.detected)
        valid_points = sum(1 for frame in self.frames for point in frame.landmarks.values() if point.valid_depth)
        slots = frames_total * len(LANDMARK_NAMES)
        return SegmentQuality(
            bag=self.bag_name,
            segment_id=self.segment.segment_id,
            label=self.segment.label,
            task_type=self.segment.task_type,
            start_s=self.segment.start_s,
            end_s=self.segment.end_s,
            frames_total=frames_total,
            frames_with_pose=frames_with_pose,
            pose_detection_rate=safe_divide(frames_with_pose, frames_total),
            valid_depth_points=valid_points,
            landmark_slots=slots,
            valid_depth_rate=safe_divide(valid_points, slots),
            keypoints_csv=str(self.keypoints_path),
            annotated_video=str(self.video_path),
        )


def run_self_test(out_dir: Optional[str]) -> None:
    target_dir = Path(out_dir).expanduser().resolve() if out_dir else Path(tempfile.mkdtemp(prefix="bodyinsp_self_test_"))
    ensure_output_dirs(target_dir)
    bag_name = "synthetic_walk"
    segment = Segment(segment_id="3", label="walk", task_type="walk", start_s=0.0, end_s=8.0)
    frames = synthetic_walk_frames(duration_s=8.0, fps=30.0)
    outputs = {segment.segment_id: InMemoryOutput(bag_name, segment, frames, target_dir)}
    metric_rows = build_metric_rows(bag_name, [segment], outputs, walk_distance_m=10.0)
    quality_rows = [quality_to_row(outputs[segment.segment_id].quality())]
    summary_rows = keypoint_summary_rows([segment], outputs)
    metrics_path = target_dir / "metrics" / f"{bag_name}_metrics.csv"
    report_path = target_dir / "reports" / f"{bag_name}_analysis.xlsx"
    write_metrics_csv(metrics_path, metric_rows)
    write_excel_report(report_path, [segment], metric_rows, quality_rows, summary_rows)
    print("Self-test complete.")
    print(f"Metrics CSV: {metrics_path}")
    print(f"Excel report: {report_path}")
    print("Synthetic metric row:")
    for metric_name in METRIC_NAMES:
        print(f"  {metric_name}: {format_float(metric_rows[0].get(metric_name), 6)}")


def build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Split a RealSense .bag by manual segments, export 3D pose keypoints, gait metrics, Excel, and annotated videos.",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter,
    )
    parser.add_argument("--bag", help="Input RealSense .bag file. Required unless --self-test is used.")
    parser.add_argument("--segments", help="CSV file with segment_id,label,task_type,start_s,end_s. Required unless --self-test is used.")
    parser.add_argument("--out", default="analysis_output", help="Output directory.")
    parser.add_argument("--score-file", default="score.csv", help="Optional name/score table used to add NP3GAIT labels to output CSV files.")
    parser.add_argument("--walk-distance-m", type=float, default=10.0, help="Walking distance used for SP_U speed calculation.")
    parser.add_argument("--depth-window", type=int, default=5, help="Odd-sized pixel window for median fallback when landmark depth is missing.")
    parser.add_argument("--model-complexity", type=int, default=1, choices=[0, 1, 2], help="MediaPipe Pose Landmarker bundle (0:lite, 1:full, 2:heavy).")
    parser.add_argument("--model-path", default=None, help="Optional local .task model path. If omitted, the selected MediaPipe bundle is cached under models/.")
    parser.add_argument("--min-detection-confidence", type=float, default=0.5, help="MediaPipe minimum detection confidence.")
    parser.add_argument("--min-tracking-confidence", type=float, default=0.5, help="MediaPipe minimum tracking confidence.")
    parser.add_argument("--video-codec", default="mp4v", help="FourCC codec for MP4 output.")
    parser.add_argument("--max-frames", type=int, default=None, help="Optional debug limit for processed frames.")
    parser.add_argument("--self-test", action="store_true", help="Run synthetic metrics and Excel export without pyrealsense2.")
    return parser


def main(argv: Optional[Sequence[str]] = None) -> int:
    parser = build_arg_parser()
    args = parser.parse_args(argv)
    try:
        if args.self_test:
            run_self_test(args.out)
            return 0
        if not args.bag:
            parser.error("--bag is required unless --self-test is used")
        if not args.segments:
            parser.error("--segments is required unless --self-test is used")
        if args.walk_distance_m <= 0:
            parser.error("--walk-distance-m must be greater than 0")
        run_analysis(args)
        return 0
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
